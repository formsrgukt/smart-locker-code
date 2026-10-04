from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

prs = Presentation()

# Layouts
# 0: Title Slide
# 8: Picture with Caption (Title, Picture, Body)
title_slide_layout = prs.slide_layouts[0]
pic_slide_layout = prs.slide_layouts[8]

# 1. Title Slide
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "SMART LOCKER"
subtitle.text = "Your Personal Digital Locker\nProject Overview & Screen Guide"

# Helper function to add a screen slide
def add_screen_slide(title_text, explanation):
    slide = prs.slides.add_slide(pic_slide_layout)
    
    # Title
    try:
        title = slide.shapes.title
        title.text = title_text
    except Exception as e:
        print(f"No title placeholder: {e}")

    # Body (Explanation)
    for shape in slide.placeholders:
        if shape.placeholder_format.type == 2: # BODY
            tf = shape.text_frame
            tf.text = explanation
            break

    # The PICTURE placeholder (type 18) remains empty, 
    # so the user sees a clickable picture icon in PowerPoint!

# 2. Landing Page
add_screen_slide(
    "Landing Page / Marketing ( / )",
    "This is the first screen users see. It features a hero section, the 'Before vs. After' Smart Locker visualization, core features, and a call-to-action to create an account or sign in."
)

# 3. Secure Login Screen
add_screen_slide(
    "Authentication ( /login )",
    "Users sign in securely using their Google Account. This screen also handles the 2-step verification (OTP sent via email) and displays a beautiful 'Secure & Private' graphic."
)

# 4. Dashboard - Home Tab
add_screen_slide(
    "Dashboard - Home",
    "The central hub of the application. It greets the user based on the time of day, shows quick action cards (Upload, Add Info, Create Collection), and displays the most recently added files."
)

# 5. Dashboard - Documents Tab
add_screen_slide(
    "Documents Tab",
    "A comprehensive view of all uploaded files. Users can search by name, filter by file type (All, Documents, Images), and perform actions like View, Share, Download, or Delete."
)

# 6. Upload / Add Document Modal
add_screen_slide(
    "Upload Document Modal",
    "Triggered from the Home or Documents tab, this modal allows users to select a file, assign it to a custom collection, and upload it securely to Firebase Storage."
)

# 7. Categories (Smart Organization)
add_screen_slide(
    "Categories Tab",
    "This screen automatically groups files into predefined categories like 'Identity', 'Education', 'Career', and 'Projects' based on their file names."
)

# 8. Collections (Custom Folders)
add_screen_slide(
    "Collections Tab",
    "Allows users to create their own custom folders (e.g., 'Tax Returns') and organize specific files inside them for easy retrieval."
)

# 9. Personal Info Management
add_screen_slide(
    "Personal Info Tab",
    "A specialized view for storing structured data like Aadhar, PAN, and Bank details securely, instead of relying on physical files. Features 1-click copy functionality."
)

# 10. Trash & Recovery
add_screen_slide(
    "Trash & Recovery Tab",
    "Acts as a recycle bin. Deleted files are moved here first, allowing users to restore them if deleted accidentally, or permanently delete them to free up space."
)

# 11. Share Document Modal
add_screen_slide(
    "Share Document Modal",
    "Allows users to generate a fast, secure CDN link for any document. Includes a clean UI with a copy-to-clipboard button that shows a green success tick when clicked."
)

output_path = "Smart_Locker_Screens_Presentation.pptx"
prs.save(output_path)
print(f"Presentation saved to {output_path}")
